# -*- coding: utf-8 -*-
"""
司机周度得分 → 司机生命周期分析报告

流程：
1. 读取「司机分数趋势底表」
2. 计算站点级 & 大区级汇总指标 → sheet「站点汇总」
3. 生成「待召回优质司机」「新司机」「流失司机」明细
4. 导出 Excel（含百分比格式、大区汇总行加粗、列宽自适应）

关键假设：底表列是 w-1 ~ w-4 四周分数，w-1 = 最近一周（本周），w-2 = 上周，w-3/w-4 更早。

| 指标 | 定义 |
|---|---|
| 本周司机人数 | w-1 分数非空 |
| 上周司机人数 | w-2 分数非空 |
| 流失司机（总体） | w-2 非空 & w-1 空 |
| 总体司机流失率 | 流失司机数 / 上周司机人数 |
| 本周优质司机 | w-1 >= 90 |
| 优质司机流失 | w-2 >= 90 & w-1 空 |
| 优质司机流失率 | 优质司机流失人数 / 上周优质司机人数 |
| 新司机 | w-1 非空，且 w-2 & w-3 & w-4 全空 |
| 新司机占比 | 新司机人数 / 本周司机人数 |
| 上周待召回优质司机 | w-2 空，且 w-3 >= 90 |
| 本周待召回优质司机 | w-2 >= 90 且 w-1 为空 |
| 本周已召回优质司机 | w-1 非空，w-2 空，且 w-3 >= 90 |
| 召回率 | 已召回优质司机人数 / 上周待召回优质司机人数 |
"""

import os

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# ============================== 配置（按需修改） ==============================
INPUT_PATH = r"E:\GOFO Driver Dashboard\History\W39司机周度得分.xlsx"      # 本周文件（W39）
SHEET_NAME = "司机分数趋势底表"

OUTPUT_DIR = r"E:\GOFO Driver Dashboard\History"
OUTPUT_NAME = "W39_Driver_Lifecycle_Report_new.xlsx"
GOOD_THRESHOLD = 90

W1, W2, W3, W4 = "w-1司机分数", "w-2司机分数", "w-3司机分数", "w-4司机分数"
COLS_OUT = ["大区", "站点", "dsp", "司机id", W1, W2, W3, W4]
# ===========================================================================


# ------------------------------ 1. 读取数据 ------------------------------
def load_data() -> pd.DataFrame:
    df = pd.read_excel(INPUT_PATH, sheet_name=SHEET_NAME)
    df.columns = [c.strip() for c in df.columns]
    print("底表形状:", df.shape)

    expected_cols = ["大区", "站点", "dsp", "司机id", W1, W2, W3, W4]
    missing = [c for c in expected_cols if c not in df.columns]
    assert not missing, f"缺少列: {missing}, 实际列名: {list(df.columns)}"

    for c in [W1, W2, W3, W4]:
        df[c] = pd.to_numeric(df[c], errors="coerce")
    return df


# ------------------------ 2. 司机粒度的辅助布尔列 ------------------------
def add_flags(df: pd.DataFrame) -> pd.DataFrame:
    df["has_w1"] = df[W1].notna()
    df["has_w2"] = df[W2].notna()
    df["has_w3"] = df[W3].notna()
    df["has_w4"] = df[W4].notna()

    df["is_good_w1"] = df["has_w1"] & (df[W1] >= GOOD_THRESHOLD)
    df["is_good_w2"] = df["has_w2"] & (df[W2] >= GOOD_THRESHOLD)
    df["is_good_w3"] = df["has_w3"] & (df[W3] >= GOOD_THRESHOLD)
    df["is_good_w4"] = df["has_w4"] & (df[W4] >= GOOD_THRESHOLD)

    # 流失：上周在岗，本周不在岗
    df["churned"] = df["has_w2"] & ~df["has_w1"]
    # 优质流失：上周是优质（w-2>=90），本周没有分数
    df["good_churned"] = df["is_good_w2"] & ~df["has_w1"]
    # 新司机：本周在岗，且过去三周从未出现过
    df["new_driver"] = df["has_w1"] & ~df["has_w2"] & ~df["has_w3"] & ~df["has_w4"]
    # 已召回优质：本周在岗，上周不在岗，但 w-3 是优质司机
    df["recalled"] = df["has_w1"] & ~df["has_w2"] & df["is_good_w3"]
    # 上周待召回优质司机：w-3 优质，w-2 没有分数（召回率的分母）
    df["pending_recall_last_week"] = df["is_good_w3"] & ~df["has_w2"]
    return df


# ------------------------------ 3. 汇总函数 ------------------------------
def safe_div(numer: pd.Series, denom: pd.Series) -> pd.Series:
    # 分母为0时返回 NaN（而不是 inf）
    return (numer / denom).replace([np.inf, -np.inf], np.nan)


def summarize_by(df_in: pd.DataFrame, group_cols: list) -> pd.DataFrame:
    agg = df_in.groupby(group_cols, as_index=False).agg(
        本周司机人数=("has_w1", "sum"),
        上周司机人数=("has_w2", "sum"),
        总体流失司机数=("churned", "sum"),
        本周优质司机人数=("is_good_w1", "sum"),
        上周优质司机人数=("is_good_w2", "sum"),
        优质司机流失人数=("good_churned", "sum"),
        新司机人数=("new_driver", "sum"),
        上周待召回优质司机人数=("pending_recall_last_week", "sum"),
        已召回优质司机人数=("recalled", "sum"),
    )

    agg["总体司机流失率"] = safe_div(agg["总体流失司机数"], agg["上周司机人数"])
    agg["优质司机流失率"] = safe_div(agg["优质司机流失人数"], agg["上周优质司机人数"])
    agg["新司机占比"] = safe_div(agg["新司机人数"], agg["本周司机人数"])
    agg["召回率"] = safe_div(agg["已召回优质司机人数"], agg["上周待召回优质司机人数"])

    agg = agg.drop(columns=["上周优质司机人数"])  # 仅中间计算用
    return agg


# ------------------- 4. 站点级汇总 + 每个大区末尾插入大区汇总行 -------------------
def build_site_report(df: pd.DataFrame) -> pd.DataFrame:
    station_summary = summarize_by(df, ["大区", "站点"])
    region_summary = summarize_by(df, ["大区"])
    region_summary["站点"] = "大区汇总"

    col_order = ["大区", "站点", "本周司机人数", "上周司机人数", "总体流失司机数", "总体司机流失率",
                 "本周优质司机人数", "优质司机流失人数", "优质司机流失率",
                 "新司机人数", "新司机占比", "上周待召回优质司机人数", "已召回优质司机人数", "召回率"]
    station_summary = station_summary[col_order]
    region_summary = region_summary[col_order]

    final_rows = []
    for region in station_summary["大区"].unique():
        final_rows.append(station_summary[station_summary["大区"] == region])
        final_rows.append(region_summary[region_summary["大区"] == region])

    return pd.concat(final_rows, ignore_index=True)


# ------------------------------ 5. 明细表 ------------------------------
def build_details(df: pd.DataFrame):
    # 待召回优质司机：上周优质、本周流失
    pending_recall = df[df["is_good_w2"] & ~df["has_w1"]][COLS_OUT].reset_index(drop=True)
    # 新司机
    new_drivers_detail = df[df["new_driver"]][COLS_OUT].reset_index(drop=True)
    # 流失司机：上周有分数、本周无分数
    churned_drivers_detail = df[df["churned"]][COLS_OUT].reset_index(drop=True)
    return pending_recall, new_drivers_detail, churned_drivers_detail


# ------------------------------ 6. 导出 Excel ------------------------------
def export_excel(site_report, pending_recall, new_drivers_detail, churned_drivers_detail) -> str:
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    output_path = os.path.join(OUTPUT_DIR, OUTPUT_NAME)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        site_report.to_excel(writer, sheet_name="站点汇总", index=False)
        pending_recall.to_excel(writer, sheet_name="本周待召回优质司机", index=False)
        new_drivers_detail.to_excel(writer, sheet_name="新司机", index=False)
        churned_drivers_detail.to_excel(writer, sheet_name="流失司机", index=False)
    print("已保存:", output_path)

    # 补充格式：百分比列、大区汇总行加粗、列宽自适应
    wb = load_workbook(output_path)
    ws = wb["站点汇总"]

    pct_cols = ["总体司机流失率", "优质司机流失率", "新司机占比", "召回率"]
    header = [cell.value for cell in ws[1]]
    pct_col_idx = [header.index(c) + 1 for c in pct_cols if c in header]
    site_cell_idx = header.index("站点") + 1

    bold_font = Font(bold=True)
    subtotal_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    for row in range(2, ws.max_row + 1):
        for col_idx in pct_col_idx:
            cell = ws.cell(row=row, column=col_idx)
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.0%"
        if ws.cell(row=row, column=site_cell_idx).value == "大区汇总":
            for col_idx in range(1, ws.max_column + 1):
                c = ws.cell(row=row, column=col_idx)
                c.font = bold_font
                c.fill = subtotal_fill

    for sheet in wb.worksheets:
        for col_cells in sheet.columns:
            length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
            col_letter = get_column_letter(col_cells[0].column)
            sheet.column_dimensions[col_letter].width = min(max(length + 2, 10), 30)

    wb.save(output_path)
    print("格式化完成:", output_path)
    return output_path


# ------------------------------ 7. 快速核对 ------------------------------
def check_regions(site_report: pd.DataFrame) -> pd.Series:
    """大区汇总的本周司机人数 是否等于各站点之和"""
    def check_region(region):
        is_total = site_report["站点"] == "大区汇总"
        in_region = site_report["大区"] == region
        stations_sum = site_report[in_region & ~is_total]["本周司机人数"].sum()
        region_total = site_report[in_region & is_total]["本周司机人数"].iloc[0]
        return stations_sum == region_total

    return pd.Series({r: check_region(r) for r in site_report["大区"].unique()})


# ================================== 主流程 ==================================
def main():
    df = load_data()
    df = add_flags(df)

    site_report = build_site_report(df)
    pending_recall, new_drivers_detail, churned_drivers_detail = build_details(df)

    export_excel(site_report, pending_recall, new_drivers_detail, churned_drivers_detail)

    print("\n大区汇总核对（True 表示大区汇总 = 各站点之和）:")
    print(check_regions(site_report))


if __name__ == "__main__":
    main()
